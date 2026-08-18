"""E2E tests for GET /api/export — full and date-ranged data export
(FUNCTIONAL_IMPROVEMENTS_PLAN.md #3, extended per DASHBOARD_COMPLETION_PLAN.md §3)."""

import os
from datetime import date

import httpx
import pytest

from workout_storage import repo
from workout_storage.db import connect

pytestmark = pytest.mark.e2e


async def _export(token, params=None):
    from workout_storage.app import app

    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://t") as c:
        return await c.get(f"/{token}/api/export", params=params or {})


async def test_full_export_shape_and_default_filename(pg_dsn, seeded):
    os.environ["DATABASE_URL"] = pg_dsn
    _, token = seeded

    r = await _export(token)
    assert r.status_code == 200
    body = r.json()
    assert body["schema_version"] == "1.0.0"
    assert len(body["exercises"]) == 10
    assert [p["id"] for p in body["programs"]] == ["prog_upper_pp"]
    assert [t["id"] for t in body["day_templates"]] == ["tpl_day_a"]
    assert len(body["sessions"]) == 1
    assert body["sessions"][0]["entries"]  # full nested shape, not a summary
    assert len(body["body_metrics"]) == 1
    assert body["coaching"]["profile"] is None  # no coach intake for this fixture
    assert body["coaching"]["goals"] == []
    assert body["coaching"]["events"] == []

    cd = r.headers["content-disposition"]
    assert cd == f'attachment; filename="aim-export-{date.today().isoformat()}.json"'


async def test_date_range_filters_sessions_and_body_metrics_only(pg_dsn, seeded):
    """The one seeded session/body_metric both fall on 2026-06; a range outside that must drop
    them while catalog/programs/day_templates — structural, not history — stay fully present."""
    os.environ["DATABASE_URL"] = pg_dsn
    _, token = seeded

    r = await _export(token, {"from": "2027-01-01", "to": "2027-01-31"})
    assert r.status_code == 200
    body = r.json()
    assert body["sessions"] == []
    assert body["body_metrics"] == []
    assert len(body["exercises"]) == 10
    assert [p["id"] for p in body["programs"]] == ["prog_upper_pp"]
    assert [t["id"] for t in body["day_templates"]] == ["tpl_day_a"]

    cd = r.headers["content-disposition"]
    assert cd == 'attachment; filename="aim-export-2027-01-01_2027-01-31.json"'


async def test_date_range_including_the_seeded_data_keeps_it(pg_dsn, seeded):
    os.environ["DATABASE_URL"] = pg_dsn
    _, token = seeded

    r = await _export(token, {"from": "2026-06-01", "to": "2026-06-30"})
    body = r.json()
    assert len(body["sessions"]) == 1
    assert len(body["body_metrics"]) == 1


async def test_coaching_section_is_never_range_filtered(pg_dsn):
    os.environ["DATABASE_URL"] = pg_dsn
    async with connect(pg_dsn) as conn:
        user = await repo.create_user(conn, name="Coached")
        await repo.ensure_coach_profile(conn, user["id"])
        await repo.update_coach_profile_fields(conn, user["id"], {"primary_goal": "strength"})
        await repo.upsert_user_goal(conn, user["id"], {"kind": "outcome", "title": "Get strong"})

    # A narrow range that couldn't possibly contain any session/body_metric for a fresh user.
    r = await _export(user["token"], {"from": "2020-01-01", "to": "2020-01-02"})
    body = r.json()
    assert body["coaching"]["profile"]["primary_goal"] == "strength"
    assert body["coaching"]["goals"][0]["title"] == "Get strong"


async def test_xlsx_export_sheets_and_rows(pg_dsn, seeded):
    """format=xlsx renders the same document as a real workbook: localized sheet names, a
    per-set row carrying the resolved exercise name (never the raw catalog id), and the
    attachment filename switching to .xlsx."""
    from io import BytesIO

    from openpyxl import load_workbook

    os.environ["DATABASE_URL"] = pg_dsn
    _, token = seeded

    r = await _export(token, {"format": "xlsx", "lang": "ru"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    cd = r.headers["content-disposition"]
    assert cd == f'attachment; filename="aim-export-{date.today().isoformat()}.xlsx"'

    wb = load_workbook(BytesIO(r.content))
    assert "Тренировки" in wb.sheetnames
    assert "Замеры тела" in wb.sheetnames
    assert "Цели" in wb.sheetnames

    sets = wb["Тренировки"]
    header = [c.value for c in sets[1]]
    assert header[0] == "Дата"
    rows = list(sets.iter_rows(min_row=2, values_only=True))
    assert rows  # the seeded session's sets are here
    # Exercise column shows a human name from the catalog, not an ex_* id.
    assert all(not str(row[2]).startswith("ex_") for row in rows)
    # Weight/reps stay numeric so the spreadsheet can actually sum them.
    assert any(isinstance(row[5], (int, float)) for row in rows)


async def test_xlsx_export_unknown_lang_falls_back_to_english(pg_dsn, seeded):
    os.environ["DATABASE_URL"] = pg_dsn
    _, token = seeded

    r = await _export(token, {"format": "xlsx", "lang": "de"})
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(r.content))
    assert "Workouts" in wb.sheetnames


async def test_token_manifest_serves_relative_start_url(pg_dsn, seeded):
    """GET /{token}/api/manifest — the token-aware PWA manifest. start_url must stay RELATIVE
    ("../app"): resolved against the manifest's own URL it becomes /{token}/app, which is the
    whole point (iOS installs from the static manifest land on "/" in a storage container with
    no saved token). "id": "/" keeps Chrome treating this as the same app as static-manifest
    installs."""
    os.environ["DATABASE_URL"] = pg_dsn
    _, token = seeded

    from workout_storage.app import app

    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://t") as c:
        r = await c.get(f"/{token}/api/manifest")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/manifest+json")
    body = r.json()
    assert body["start_url"] == "../app"
    assert body["id"] == "/"
    assert body["scope"] == "/"
    assert body["display"] == "standalone"
    assert {i["src"] for i in body["icons"]} >= {"/pwa-192.png", "/pwa-512.png"}


async def test_partial_range_only_from(pg_dsn, seeded):
    os.environ["DATABASE_URL"] = pg_dsn
    _, token = seeded

    r = await _export(token, {"from": "2026-06-06"})
    body = r.json()
    assert len(body["sessions"]) == 1
    assert (
        r.headers["content-disposition"] == 'attachment; filename="aim-export-2026-06-06_now.json"'
    )
