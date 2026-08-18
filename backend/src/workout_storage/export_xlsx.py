"""Excel rendering of the export document (GET /api/export?format=xlsx).

The JSON export is the machine-readable backup that import_document can restore; this module is
the human-readable counterpart for people who open their data in Excel/Numbers/Sheets and have
never heard of JSON. It renders the same export_document() dict — never its own queries — so the
two formats can't drift apart in content.

Sheet layout is flat-table-per-topic (one row per set / session / measurement / goal / program
item): the shapes spreadsheet users actually sort, filter and chart. Nested JSON-only detail
(custom_fields, HR zones, coach events) stays JSON-only on purpose.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Header/sheet strings per app locale. Keys mirror the frontend's supported languages; anything
# unknown falls back to English, same as the app's i18n.
_L10N: dict[str, dict[str, str]] = {
    "en": {
        "sheet_sets": "Workouts",
        "sheet_sessions": "Workout summary",
        "sheet_cardio": "Cardio",
        "sheet_body": "Body measurements",
        "sheet_goals": "Goals",
        "sheet_program": "Program",
        "date": "Date",
        "day": "Day",
        "exercise": "Exercise",
        "set_number": "Set #",
        "set_type": "Set type",
        "weight_kg": "Weight (kg)",
        "reps": "Reps",
        "rir": "RIR",
        "rpe": "RPE",
        "duration_sec": "Duration (sec)",
        "distance_m": "Distance (m)",
        "notes": "Notes",
        "status": "Status",
        "duration_min": "Duration (min)",
        "session_rpe": "Session RPE",
        "bodyweight_kg": "Body weight (kg)",
        "volume_kg": "Volume (kg)",
        "tags": "Tags",
        "cardio_type": "Type",
        "avg_hr": "Avg HR",
        "calories": "Calories",
        "body_fat_pct": "Body fat (%)",
        "goal": "Goal",
        "goal_type": "Type",
        "metric": "Metric",
        "baseline": "Starting value",
        "target": "Target value",
        "unit": "Unit",
        "deadline": "Deadline",
        "review_date": "Review date",
        "created": "Created",
        "program": "Program",
        "block": "Block",
        "target_sets": "Sets",
        "target_reps": "Reps",
        "target_weight_kg": "Target weight (kg)",
        "rest_sec": "Rest (sec)",
    },
    "ru": {
        "sheet_sets": "Тренировки",
        "sheet_sessions": "Сводка тренировок",
        "sheet_cardio": "Кардио",
        "sheet_body": "Замеры тела",
        "sheet_goals": "Цели",
        "sheet_program": "Программа",
        "date": "Дата",
        "day": "День",
        "exercise": "Упражнение",
        "set_number": "Подход",
        "set_type": "Тип подхода",
        "weight_kg": "Вес (кг)",
        "reps": "Повторы",
        "rir": "RIR",
        "rpe": "RPE",
        "duration_sec": "Время (сек)",
        "distance_m": "Дистанция (м)",
        "notes": "Заметки",
        "status": "Статус",
        "duration_min": "Длительность (мин)",
        "session_rpe": "RPE тренировки",
        "bodyweight_kg": "Вес тела (кг)",
        "volume_kg": "Объём (кг)",
        "tags": "Теги",
        "cardio_type": "Тип",
        "avg_hr": "Ср. пульс",
        "calories": "Калории",
        "body_fat_pct": "Жир (%)",
        "goal": "Цель",
        "goal_type": "Тип",
        "metric": "Метрика",
        "baseline": "Стартовое значение",
        "target": "Целевое значение",
        "unit": "Единица",
        "deadline": "Дедлайн",
        "review_date": "Дата пересмотра",
        "created": "Создана",
        "program": "Программа",
        "block": "Блок",
        "target_sets": "Подходы",
        "target_reps": "Повторы",
        "target_weight_kg": "Целевой вес (кг)",
        "rest_sec": "Отдых (сек)",
    },
    "pt": {
        "sheet_sets": "Treinos",
        "sheet_sessions": "Resumo das sessões",
        "sheet_cardio": "Cardio",
        "sheet_body": "Medidas corporais",
        "sheet_goals": "Objetivos",
        "sheet_program": "Programa",
        "date": "Data",
        "day": "Dia",
        "exercise": "Exercício",
        "set_number": "Série",
        "set_type": "Tipo de série",
        "weight_kg": "Peso (kg)",
        "reps": "Reps",
        "rir": "RIR",
        "rpe": "RPE",
        "duration_sec": "Duração (seg)",
        "distance_m": "Distância (m)",
        "notes": "Notas",
        "status": "Status",
        "duration_min": "Duração (min)",
        "session_rpe": "RPE da sessão",
        "bodyweight_kg": "Peso corporal (kg)",
        "volume_kg": "Volume (kg)",
        "tags": "Etiquetas",
        "cardio_type": "Tipo",
        "avg_hr": "FC média",
        "calories": "Calorias",
        "body_fat_pct": "Gordura corporal (%)",
        "goal": "Objetivo",
        "goal_type": "Tipo",
        "metric": "Métrica",
        "baseline": "Valor inicial",
        "target": "Valor alvo",
        "unit": "Unidade",
        "deadline": "Prazo",
        "review_date": "Data de revisão",
        "created": "Criado",
        "program": "Programa",
        "block": "Bloco",
        "target_sets": "Séries",
        "target_reps": "Reps",
        "target_weight_kg": "Peso alvo (kg)",
        "rest_sec": "Descanso (seg)",
    },
    "es": {
        "sheet_sets": "Entrenamientos",
        "sheet_sessions": "Resumen de sesiones",
        "sheet_cardio": "Cardio",
        "sheet_body": "Medidas corporales",
        "sheet_goals": "Objetivos",
        "sheet_program": "Programa",
        "date": "Fecha",
        "day": "Día",
        "exercise": "Ejercicio",
        "set_number": "Serie",
        "set_type": "Tipo de serie",
        "weight_kg": "Peso (kg)",
        "reps": "Reps",
        "rir": "RIR",
        "rpe": "RPE",
        "duration_sec": "Duración (seg)",
        "distance_m": "Distancia (m)",
        "notes": "Notas",
        "status": "Estado",
        "duration_min": "Duración (min)",
        "session_rpe": "RPE de la sesión",
        "bodyweight_kg": "Peso corporal (kg)",
        "volume_kg": "Volumen (kg)",
        "tags": "Etiquetas",
        "cardio_type": "Tipo",
        "avg_hr": "FC media",
        "calories": "Calorías",
        "body_fat_pct": "Grasa corporal (%)",
        "goal": "Objetivo",
        "goal_type": "Tipo",
        "metric": "Métrica",
        "baseline": "Valor inicial",
        "target": "Valor objetivo",
        "unit": "Unidad",
        "deadline": "Fecha límite",
        "review_date": "Fecha de revisión",
        "created": "Creado",
        "program": "Programa",
        "block": "Bloque",
        "target_sets": "Series",
        "target_reps": "Reps",
        "target_weight_kg": "Peso objetivo (kg)",
        "rest_sec": "Descanso (seg)",
    },
    "fr": {
        "sheet_sets": "Entraînements",
        "sheet_sessions": "Résumé des séances",
        "sheet_cardio": "Cardio",
        "sheet_body": "Mesures corporelles",
        "sheet_goals": "Objectifs",
        "sheet_program": "Programme",
        "date": "Date",
        "day": "Jour",
        "exercise": "Exercice",
        "set_number": "Série",
        "set_type": "Type de série",
        "weight_kg": "Poids (kg)",
        "reps": "Réps",
        "rir": "RIR",
        "rpe": "RPE",
        "duration_sec": "Durée (sec)",
        "distance_m": "Distance (m)",
        "notes": "Notes",
        "status": "Statut",
        "duration_min": "Durée (min)",
        "session_rpe": "RPE de la séance",
        "bodyweight_kg": "Poids corporel (kg)",
        "volume_kg": "Volume (kg)",
        "tags": "Tags",
        "cardio_type": "Type",
        "avg_hr": "FC moyenne",
        "calories": "Calories",
        "body_fat_pct": "Masse grasse (%)",
        "goal": "Objectif",
        "goal_type": "Type",
        "metric": "Métrique",
        "baseline": "Valeur de départ",
        "target": "Valeur cible",
        "unit": "Unité",
        "deadline": "Échéance",
        "review_date": "Date de révision",
        "created": "Créé",
        "program": "Programme",
        "block": "Bloc",
        "target_sets": "Séries",
        "target_reps": "Réps",
        "target_weight_kg": "Poids cible (kg)",
        "rest_sec": "Repos (sec)",
    },
}


def _exercise_names(doc: dict[str, Any]) -> dict[str, str]:
    return {e["id"]: e["name"] for e in doc.get("exercises", []) if e.get("id") and e.get("name")}


def _exercise_label(exercise_id: str | None, names: dict[str, str]) -> str:
    if not exercise_id:
        return ""
    if exercise_id in names:
        return names[exercise_id]
    # Same last-resort de-slug the app uses for orphaned catalog ids (ex_goblet_squat →
    # "goblet squat") — a readable guess beats leaking a raw id into a spreadsheet.
    return exercise_id.removeprefix("ex_").replace("_", " ")


def _rep_range(target_reps: dict[str, Any] | None) -> str:
    if not target_reps:
        return ""
    lo, hi = target_reps.get("min"), target_reps.get("max")
    if lo is not None and hi is not None:
        return str(lo) if lo == hi else f"{lo}-{hi}"
    return str(lo if lo is not None else hi if hi is not None else "")


def _day(value: str | None) -> str:
    # created_at style timestamps → keep just the date part; date strings pass through.
    return (value or "")[:10]


def _add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    ws: Worksheet = wb.create_sheet(title=title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for idx, header in enumerate(headers, start=1):
        width = max(
            len(str(header)),
            *(len(str(r[idx - 1])) for r in rows if r[idx - 1] is not None),
            0,
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 8), 40)


def _set_rows(doc: dict[str, Any], names: dict[str, str]) -> list[list[Any]]:
    rows = []
    for session in doc.get("sessions", []):
        for entry in session.get("entries", []):
            exercise = _exercise_label(entry.get("exercise_id"), names)
            for s in entry.get("sets", []):
                rows.append(
                    [
                        session.get("date"),
                        session.get("day_label"),
                        exercise,
                        s.get("set_number"),
                        s.get("type"),
                        s.get("weight_kg"),
                        s.get("reps"),
                        s.get("rir"),
                        s.get("rpe"),
                        s.get("duration_sec"),
                        s.get("distance_m"),
                        s.get("notes") or entry.get("notes"),
                    ]
                )
    return rows


def _session_rows(doc: dict[str, Any]) -> list[list[Any]]:
    rows = []
    for session in doc.get("sessions", []):
        volume = sum(
            (e.get("total_volume_kg") or 0)
            for e in session.get("entries", [])
            if e.get("total_volume_kg") is not None
        )
        duration = session.get("duration_sec")
        rows.append(
            [
                session.get("date"),
                session.get("day_label"),
                session.get("status"),
                round(duration / 60) if duration else None,
                session.get("session_rpe"),
                session.get("bodyweight_kg"),
                round(volume) if volume else None,
                ", ".join(session.get("tags") or []),
                session.get("notes"),
            ]
        )
    return rows


def _cardio_rows(doc: dict[str, Any]) -> list[list[Any]]:
    rows = []
    for session in doc.get("sessions", []):
        for c in session.get("cardio", []):
            rows.append(
                [
                    session.get("date"),
                    c.get("type"),
                    c.get("distance_m"),
                    c.get("duration_sec"),
                    c.get("avg_hr"),
                    c.get("calories"),
                    c.get("notes"),
                ]
            )
    return rows


def _body_rows(doc: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    metrics = doc.get("body_metrics", [])
    extra_keys = sorted({k for m in metrics for k in (m.get("measurements") or {})})
    rows = []
    for m in metrics:
        measurements = m.get("measurements") or {}
        rows.append(
            [m.get("date"), m.get("bodyweight_kg"), m.get("body_fat_pct")]
            + [measurements.get(k) for k in extra_keys]
            + [m.get("notes")]
        )
    return extra_keys, rows


def _goal_rows(doc: dict[str, Any]) -> list[list[Any]]:
    rows = []
    for g in doc.get("coaching", {}).get("goals", []):
        target = g.get("target") or {}
        rows.append(
            [
                g.get("title"),
                g.get("goal_type") or target.get("goal_type"),
                g.get("status"),
                target.get("metric"),
                target.get("baseline_value"),
                target.get("value"),
                target.get("unit"),
                target.get("deadline"),
                g.get("review_date"),
                _day(g.get("created_at")),
                g.get("notes"),
            ]
        )
    return rows


def _program_rows(doc: dict[str, Any], names: dict[str, str]) -> list[list[Any]]:
    programs = {p["id"]: p for p in doc.get("programs", []) if p.get("id")}
    rows = []
    for tpl in doc.get("day_templates", []):
        program = programs.get(tpl.get("program_id") or "", {})
        for block in tpl.get("blocks", []):
            for item in block.get("items", []):
                rows.append(
                    [
                        program.get("name"),
                        tpl.get("name"),
                        block.get("label") or block.get("type"),
                        _exercise_label(item.get("exercise_id"), names),
                        item.get("target_sets"),
                        _rep_range(item.get("target_reps")),
                        item.get("target_weight_kg"),
                        item.get("rest_sec"),
                        item.get("notes"),
                    ]
                )
    return rows


def render_xlsx(doc: dict[str, Any], lang: str = "en") -> bytes:
    """Render an export_document() dict as an .xlsx file, headers in the given app language."""
    t = _L10N.get(lang, _L10N["en"])
    names = _exercise_names(doc)

    wb = Workbook()
    # openpyxl always creates one default sheet; build our own and drop it at the end.
    default = wb.active

    _add_sheet(
        wb,
        t["sheet_sets"],
        [
            t["date"],
            t["day"],
            t["exercise"],
            t["set_number"],
            t["set_type"],
            t["weight_kg"],
            t["reps"],
            t["rir"],
            t["rpe"],
            t["duration_sec"],
            t["distance_m"],
            t["notes"],
        ],
        _set_rows(doc, names),
    )
    _add_sheet(
        wb,
        t["sheet_sessions"],
        [
            t["date"],
            t["day"],
            t["status"],
            t["duration_min"],
            t["session_rpe"],
            t["bodyweight_kg"],
            t["volume_kg"],
            t["tags"],
            t["notes"],
        ],
        _session_rows(doc),
    )
    cardio_rows = _cardio_rows(doc)
    if cardio_rows:
        _add_sheet(
            wb,
            t["sheet_cardio"],
            [
                t["date"],
                t["cardio_type"],
                t["distance_m"],
                t["duration_sec"],
                t["avg_hr"],
                t["calories"],
                t["notes"],
            ],
            cardio_rows,
        )
    extra_keys, body_rows = _body_rows(doc)
    _add_sheet(
        wb,
        t["sheet_body"],
        [t["date"], t["bodyweight_kg"], t["body_fat_pct"], *extra_keys, t["notes"]],
        body_rows,
    )
    _add_sheet(
        wb,
        t["sheet_goals"],
        [
            t["goal"],
            t["goal_type"],
            t["status"],
            t["metric"],
            t["baseline"],
            t["target"],
            t["unit"],
            t["deadline"],
            t["review_date"],
            t["created"],
            t["notes"],
        ],
        _goal_rows(doc),
    )
    program_rows = _program_rows(doc, names)
    if program_rows:
        _add_sheet(
            wb,
            t["sheet_program"],
            [
                t["program"],
                t["day"],
                t["block"],
                t["exercise"],
                t["target_sets"],
                t["target_reps"],
                t["target_weight_kg"],
                t["rest_sec"],
                t["notes"],
            ],
            program_rows,
        )

    if default is not None:
        wb.remove(default)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
